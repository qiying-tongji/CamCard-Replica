# -*- coding: utf-8 -*-
"""Generate fictional CamCard test cards and the ground-truth workbook.

The course document asks the teacher to supply 30–50 shared cards. This copy of
the assignment does not include that set. These cards are synthetic stand-ins
until an official set arrives. They must not replace a teacher-provided set.
"""

from __future__ import annotations

import os
from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageFont
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.abspath(__file__))
IMAGE_DIR = os.path.join(ROOT, "test_card_images")
XLSX_PATH = os.path.join(ROOT, "05_test_cards.xlsx")

FONT_REGULAR = r"C:\Windows\Fonts\msyh.ttc"
FONT_BOLD = r"C:\Windows\Fonts\msyhbd.ttc"


def card(**kwargs):
    base = {
        "name": "",
        "company": "",
        "department": "",
        "jobTitle": "",
        "mobile": "",
        "phone": "",
        "fax": "",
        "email": "",
        "website": "",
        "address": "",
        "postal_code": "",
        "company_en": "",
        "department_en": "",
        "logo_text": "",
        "lang": "en",
        "layout": "横版",
        "capture_condition": "正面清晰",
        "theme": "light",
    }
    base.update(kwargs)
    return base


CARDS = [
    card(
        card_id="C001", category="英文名片", lang="en", theme="light",
        capture_condition="正面清晰",
        name="Elena Marsh", company="Northwind Labs", jobTitle="Product Manager",
        mobile="+1-555-0101", phone="+1-555-0181",
        email="elena.marsh@example.com", website="www.northwind.example",
        address="18 Harbor Road, Qinglan", postal_code="990018",
        notes="英文横版，字段齐全，含网站。适合作为 T07、T08 的标准样例。",
    ),
    card(
        card_id="C002", category="英文名片", lang="en", theme="paper",
        capture_condition="倾斜",
        name="Hugo Bennett", company="Brightpath Analytics", jobTitle="Data Scientist",
        mobile="+1-555-0102", phone="+1-555-0182",
        email="hugo.bennett@example.com", website="www.brightpath.example",
        address="42 King Street, Qinglan", postal_code="990042",
        notes="名片放在浅色桌面上并倾斜，用于 T06 边界检测。文字本身仍然完整。",
    ),
    card(
        card_id="C003", category="英文名片", lang="en", theme="light",
        capture_condition="深色背景",
        name="Priya Nair", company="Lumen Forge", jobTitle="Design Lead",
        mobile="+1-555-0103",
        email="priya.nair@example.com", website="www.lumenforge.example",
        address="9 Lantern Lane, Qinglan",
        notes="深色桌面背景。没有办公室电话。网站存在。",
    ),
    card(
        card_id="C004", category="英文名片", lang="en", theme="paper",
        capture_condition="浅色背景",
        name="Omar Haddad", company="Cedar Grid", department="Platform Group",
        jobTitle="Solutions Architect",
        mobile="+1-555-0104", phone="+1-555-0184",
        email="omar.haddad@example.com", website="www.cedargrid.example",
        address="77 Cedar Avenue, Qinglan", postal_code="990077",
        notes="浅色背景。部门与公司分开，部门不是公司名。",
    ),
    card(
        card_id="C005", category="英文名片", lang="en", theme="light",
        capture_condition="带阴影",
        name="Sofia Almeida", company="Harbor & Pine", jobTitle="Account Director",
        mobile="+1-555-0105",
        email="sofia.almeida@example.com", website="www.harborpine.example",
        address="5 Pine Wharf, Qinglan",
        notes="名片带阴影。公司名中的 & 必须保留。没有办公室电话。",
    ),
    card(
        card_id="C006", category="英文名片", lang="en", theme="paper",
        capture_condition="边缘遮挡",
        name="Noah Keller", company="Kindred Signal", jobTitle="Marketing Manager",
        mobile="+1-555-0106", phone="+1-555-0186",
        email="noah.keller@example.com", website="www.kindredsignal.example",
        address="12 Signal Court, Qinglan",
        notes="遮挡只压住名片外圈空白，姓名和联系字段仍可见。用于边界检测，不是字段缺失。",
    ),
    card(
        card_id="C007", category="英文名片", lang="en", theme="light",
        capture_condition="正面清晰",
        name="Grace Okonkwo", company="Fieldnote Studio", jobTitle="Editor",
        mobile="+1-555-0107",
        email="grace.okonkwo@example.com",
        address="30 Notebook Road, Qinglan",
        notes="T09 主样例。名片上没有网站。website 必须为空，不得补全。",
    ),
    card(
        card_id="C008", category="中文名片", lang="zh", theme="paper",
        capture_condition="正面清晰",
        name="林晓舟", company="青岚科技", department="产品部", jobTitle="产品经理",
        mobile="+86 138-0000-1008", phone="+86 21-0000-2008",
        email="xiaozhou.lin@example.com", website="www.qinglan-tech.example",
        address="青岚市海港路18号", postal_code="990018",
        notes="中文横版，字段齐全。部门「产品部」不得写入公司或职位。",
    ),
    card(
        card_id="C009", category="中文名片", lang="zh", theme="light",
        capture_condition="倾斜",
        name="赵清和", company="远航物流", jobTitle="运营总监",
        mobile="+86 138-0000-1009", phone="+86 21-0000-2009",
        email="qinghe.zhao@example.com", website="www.yuanhang.example",
        address="澄湾市码头路6号",
        notes="中文名片，倾斜放置在桌面上。",
    ),
    card(
        card_id="C010", category="中文名片", lang="zh", theme="paper",
        capture_condition="深色背景",
        name="孙若溪", company="澄光教育", department="课程中心", jobTitle="课程顾问",
        mobile="+86 138-0000-1010",
        email="ruoxi.sun@example.com",
        address="青岚市书院路21号",
        notes="深色背景，没有网站，没有办公室电话。",
    ),
    card(
        card_id="C011", category="中文名片", lang="zh", theme="light",
        capture_condition="正面清晰",
        name="吴景行", company="北窗建筑", jobTitle="项目建筑师",
        mobile="+86 138-0000-1011", phone="+86 21-0000-2011",
        email="jingxing.wu@example.com",
        notes="没有网站，也没有地址。不得根据公司名编造网址或街道。",
    ),
    card(
        card_id="C012", category="中文名片", lang="zh", theme="paper",
        capture_condition="小字号",
        name="郑曼宁", company="拾穗文化", jobTitle="品牌主理人",
        mobile="+86 138-0000-1012",
        email="manning.zheng@example.com", website="www.shisui.example",
        address="青岚市文昌巷3号",
        notes="字号较小，但字段完整且含网站。",
    ),
    card(
        card_id="C013", category="中文名片", lang="zh", theme="light",
        capture_condition="正面清晰",
        name="何予安", company="津渡咨询", jobTitle="分析师",
        mobile="+86 138-0000-1013", phone="+86 21-0000-2013",
        email="yuan.he@example.com", website="www.jindu.example",
        address="澄湾市渡口街15号",
        notes="中文标准横版。邮箱本地部分是 yuan.he，姓名是何予安，二者不要互换。",
    ),
    card(
        card_id="C014", category="中文名片", lang="zh", theme="paper",
        capture_condition="复杂桌面背景",
        name="马听澜", company="石桥制造", department="质量部", jobTitle="质量主管",
        mobile="+86 138-0000-1014", phone="+86 21-0000-2014", fax="+86 21-0000-3014",
        email="tinglan.ma@example.com", website="www.shiqiao.example",
        address="青岚市石桥路88号", postal_code="990088",
        notes="复杂桌面背景。同时有手机、办公室电话和传真。",
    ),
    card(
        card_id="C015", category="中英双语", lang="bi", theme="light",
        capture_condition="正面清晰",
        name="周启年", company="青岚大学", company_en="Qinglan University",
        department="计算机学院", department_en="School of Computer Science",
        jobTitle="Professor",
        mobile="+86 138-0000-1015",
        email="qinian.zhou@example.com",
        address="青岚市大学路1号",
        notes="中英混排。姓名只保留「周启年」，不得把 Professor 并入姓名。公司记「青岚大学」；Qinglan University 是同一机构的英文名，不得另记第二家公司。没有网站。",
    ),
    card(
        card_id="C016", category="中英双语", lang="bi", theme="paper",
        capture_condition="正面清晰",
        name="陈见微", company="岚桥设计", company_en="Lanqiao Design",
        jobTitle="Design Director",
        mobile="+86 138-0000-1016",
        email="jianwei.chen@example.com",
        address="青岚市设计路9号",
        notes="无网站。Lanqiao Design 与「岚桥设计」是同一公司。",
    ),
    card(
        card_id="C017", category="中英双语", lang="bi", theme="light",
        capture_condition="倾斜",
        name="许知夏", company="南风医疗", company_en="Nanfeng Health",
        department="研究部", jobTitle="Research Associate",
        mobile="+86 138-0000-1017", phone="+86 21-0000-2017",
        email="zhixia.xu@example.com", website="www.nanfeng.example",
        address="澄湾市健康路12号",
        notes="双语名片，倾斜放置。职位使用英文，不得并入姓名。",
    ),
    card(
        card_id="C018", category="中英双语", lang="bi", theme="paper",
        capture_condition="浅色背景",
        name="罗景深", company="白塔软件", company_en="Baita Software",
        jobTitle="Engineering Manager",
        mobile="+86 138-0000-1018",
        email="jingshen.luo@example.com", website="www.baita.example",
        address="青岚市软件园路7号",
        notes="浅色背景上的双语名片。",
    ),
    card(
        card_id="C019", category="中英双语", lang="bi", theme="light",
        capture_condition="正面清晰",
        name="邓书白", company="合页资本", company_en="Heye Capital",
        jobTitle="Associate",
        mobile="+1-555-0119",
        email="shubai.deng@example.com",
        notes="无网站、无地址。职位 Associate 不得并入姓名，也不得当成公司。",
    ),
    card(
        card_id="C020", category="中英双语", lang="bi", theme="paper",
        capture_condition="带阴影",
        name="冯晚晴", company="橘园媒体", company_en="Juyuan Media",
        jobTitle="Producer",
        mobile="+86 138-0000-1020", phone="+86 21-0000-2020",
        email="wanqing.feng@example.com", website="www.juyuan.example",
        address="青岚市橘园路4号",
        notes="带阴影的双语名片。",
    ),
    card(
        card_id="C021", category="竖版/特殊布局", lang="zh", theme="paper",
        layout="竖版", capture_condition="正面清晰",
        name="沈予白", company="墨白工作室", jobTitle="独立设计师",
        mobile="+86 138-0000-1021",
        email="yubai.shen@example.com",
        address="青岚市墨香里2号",
        notes="竖版中文。没有网站。",
    ),
    card(
        card_id="C022", category="竖版/特殊布局", lang="en", theme="light",
        layout="竖版", capture_condition="正面清晰",
        name="Caleb Orth", company="Orth Meridian", jobTitle="Field Engineer",
        mobile="+1-555-0122", phone="+1-555-0172",
        email="caleb.orth@example.com", website="www.orthmeridian.example",
        address="14 Meridian Way, Qinglan",
        notes="竖版英文。公司 Orth Meridian 与姓 Orth 相邻。姓名只能是 Caleb Orth，不能把公司整句写入姓名。",
    ),
    card(
        card_id="C023", category="竖版/特殊布局", lang="bi", theme="paper",
        layout="竖版", capture_condition="正面清晰",
        name="韩澄", company="澄川实验室", company_en="Chengchuan Lab",
        jobTitle="Lab Manager",
        mobile="+86 138-0000-1023",
        email="cheng.han@example.com", website="www.chengchuan.example",
        address="澄湾市实验路5号",
        notes="竖版双语。",
    ),
    card(
        card_id="C024", category="竖版/特殊布局", lang="zh", theme="light",
        layout="特殊分栏", capture_condition="正面清晰",
        name="梁小满", company="满仓零售", department="门店", jobTitle="店长",
        mobile="+86 138-0000-1024", phone="+86 21-0000-2024",
        email="xiaoman.liang@example.com",
        address="青岚市集市路27号",
        notes="横版左右分栏：左栏是身份，右栏是联系方式。没有网站。",
    ),
    card(
        card_id="C025", category="多电话/多Email", lang="en", theme="paper",
        capture_condition="正面清晰",
        name="Irene Cho", company="Copperline Freight", jobTitle="Logistics Lead",
        mobile="+1-555-0125", phone="+1-555-0185", fax="+1-555-0195",
        email="irene.cho@example.com",
        address="410 Freight Road, Qinglan",
        notes="T10 主样例。同时有 Mobile、Office、Fax。没有网站。三个号码必须分字段保存，传真不得写入手机。",
    ),
    card(
        card_id="C026", category="多电话/多Email", lang="zh", theme="light",
        capture_condition="正面清晰",
        name="潘星河", company="星河会展", jobTitle="会务主管",
        mobile="+86 138-0000-1026", phone="+86 21-0000-2026", fax="+86 21-0000-3026",
        email="xinghe.pan@example.com", website="www.xinghe.example",
        address="青岚市会展路8号",
        notes="中文三号码样例：手机、电话、传真同时存在，且含网站。",
    ),
    card(
        card_id="C027", category="多电话/多Email", lang="en", theme="paper",
        capture_condition="正面清晰",
        name="Samuel Iqbal", company="Two Rivers Legal", jobTitle="Counsel",
        mobile="+1-555-0127", phone="+1-555-0187",
        email="samuel.iqbal@example.com; s.iqbal@tworivers.example",
        website="www.tworivers.example",
        address="2 River Street, Qinglan",
        notes="两个邮箱都要保留，表格里用分号加空格分隔。不得只留一个，不得把邮箱本地部分当成姓名。",
    ),
    card(
        card_id="C028", category="多电话/多Email", lang="zh", theme="light",
        capture_condition="正面清晰",
        name="叶知秋", company="秋水翻译", jobTitle="译员",
        mobile="+86 138-0000-1028", phone="+86 21-0000-2028",
        email="zhiqiu.ye@example.com; ye.zhiqiu@example.com",
        notes="两个邮箱、手机和办公室电话。没有网站，没有地址，没有传真。",
    ),
    card(
        card_id="C029", category="字段缺失", lang="en", theme="paper",
        capture_condition="正面清晰",
        name="Nora Blake", company="Paper Lantern Press",
        mobile="+1-555-0129",
        email="nora.blake@example.com",
        address="48 Paper Street, Qinglan",
        notes="没有职位，没有网站。不得根据出版社猜测职位或网址。",
    ),
    card(
        card_id="C030", category="字段缺失", lang="zh", theme="light",
        capture_condition="正面清晰",
        name="", company="岚桥设计事务所",
        phone="+86 21-0000-2030",
        email="studio@lanqiao.example",
        address="青岚市设计路9号",
        notes="没有个人姓名，也没有网站。不得把公司名写入姓名。",
    ),
    card(
        card_id="C031", category="字段缺失", lang="zh", theme="paper",
        capture_condition="正面清晰",
        name="曹立衡", company="立衡会计",
        mobile="+86 138-0000-1031",
        email="liheng.cao@example.com",
        address="澄湾市算盘街11号",
        notes="没有职位，没有网站。公司与姓名共享「立衡」二字，仍须分开。",
    ),
    card(
        card_id="C032", category="字段缺失", lang="zh", theme="light",
        capture_condition="正面清晰",
        name="丁小山",
        mobile="+86 138-0000-1032",
        notes="只有姓名和手机。公司、职位、办公室电话、传真、邮箱、网站、地址都必须为空。",
    ),
    card(
        card_id="C033", category="字段歧义", lang="en", theme="paper",
        capture_condition="正面清晰",
        name="Jordan Lee", company="Lee Consulting", jobTitle="Founder",
        mobile="+1-555-0133",
        email="jordan.lee@example.com",
        notes="人名与公司都含 Lee。姓名不得变成 Lee Consulting，公司不得变成 Jordan Lee。没有网站和地址。",
    ),
    card(
        card_id="C034", category="字段歧义", lang="en", theme="light",
        capture_condition="正面清晰",
        name="Mina Cole",
        mobile="+1-555-0134",
        email="mina.cole@example.com",
        address="8 Apple Street, Qinglan",
        notes="Apple 只出现在地址里。公司必须为空，网站必须为空，不得把 Apple 写成公司或网址。",
    ),
    card(
        card_id="C035", category="拍摄质量较差", lang="zh", theme="paper",
        capture_condition="光线不足且文字较小",
        name="袁慢", company="慢递印刷", jobTitle="印刷顾问",
        mobile="+86 138-0000-1035", phone="+86 21-0000-2035",
        email="man.yuan@example.com", website="www.mandi.example",
        address="青岚市印刷巷13号",
        notes="模拟光线不足和小字号。正确答案仍然是这些字段，画质不能作为编造额外字段的理由。",
    ),
    card(
        card_id="C036", category="拍摄质量较差", lang="en", theme="light",
        capture_condition="模糊低对比",
        name="Theo Marlow", company="Glasshouse Survey", jobTitle="Surveyor",
        mobile="+1-555-0136",
        email="theo.marlow@example.com", website="www.glasshouse.example",
        address="16 Glass Road, Qinglan",
        notes="模拟轻微模糊和低对比。网站存在，不得漏当成缺失后再编造另一个网址。",
    ),
    card(
        card_id="C037", category="拍摄质量较差", lang="zh", theme="paper",
        capture_condition="倾斜加阴影加小字",
        name="蒋南枝", company="南枝花艺", jobTitle="花艺师",
        mobile="+86 138-0000-1037",
        email="nanzhi.jiang@example.com",
        address="澄湾市花市口2号",
        notes="倾斜、阴影、小字同时出现。没有网站。画质差不能成为编造网站的理由。",
    ),
    card(
        card_id="C038", category="复杂Logo/背景", lang="zh", theme="light",
        capture_condition="复杂Logo",
        name="曲晓棠", company="棠口食品", jobTitle="品牌经理",
        mobile="+86 138-0000-1038", phone="+86 21-0000-2038",
        email="xiaotang.qu@example.com", website="www.tangkou.example",
        address="青岚市食品园路19号",
        logo_text="TK FOODS",
        notes="Logo 文字是 TK FOODS，不是姓名，也不是公司。姓名是曲晓棠，公司是棠口食品。",
    ),
    card(
        card_id="C039", category="复杂Logo/背景", lang="en", theme="paper",
        capture_condition="复杂背景",
        name="Helen Cho", company="Vivid Parcel", jobTitle="Operations Lead",
        mobile="+1-555-0139", phone="+1-555-0189",
        email="helen.cho@example.com", website="www.vividparcel.example",
        address="90 Parcel Yard, Qinglan",
        logo_text="VP",
        notes="背景杂乱，Logo 字母是 VP，不是姓名。此人与 C025 Irene Cho 不是同一个人。",
    ),
    card(
        card_id="C040", category="复杂Logo/背景", lang="bi", theme="light",
        capture_condition="复杂桌面背景",
        name="柏年", company="柏年实验室", company_en="Bainian Lab",
        jobTitle="研究员",
        mobile="+86 138-0000-1040",
        email="nian.bai@example.com", website="www.bainianlab.example",
        address="青岚市柏树路10号",
        logo_text="BN LAB",
        notes="姓名「柏年」与公司「柏年实验室」共享前两字。公司不得只剩「柏年」，姓名不得变成整句公司名。Logo「BN LAB」不是姓名。英文 Bainian Lab 是同一机构。",
    ),
]


THEMES = {
    "light": {"bg": (247, 244, 239), "ink": (28, 28, 28), "muted": (70, 74, 80), "accent": (31, 78, 121)},
    "paper": {"bg": (255, 252, 247), "ink": (36, 32, 28), "muted": (78, 68, 56), "accent": (140, 74, 43)},
}


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    path = FONT_BOLD if bold else FONT_REGULAR
    return ImageFont.truetype(path, size)


def wrap_text(draw: ImageDraw.ImageDraw, text: str, face: ImageFont.FreeTypeFont, max_width: int) -> list[str]:
    if not text:
        return []
    lines = []
    current = ""
    for ch in text:
        trial = current + ch
        if draw.textlength(trial, font=face) <= max_width or not current:
            current = trial
        else:
            lines.append(current)
            current = ch
    if current:
        lines.append(current)
    return lines


def labels_for(lang: str) -> dict[str, str]:
    if lang == "zh":
        return {"mobile": "手机", "phone": "电话", "fax": "传真", "email": "邮箱", "website": "网址", "address": "地址", "postal": "邮编"}
    if lang == "bi":
        return {
            "mobile": "手机 Mobile",
            "phone": "电话 Office",
            "fax": "传真 Fax",
            "email": "邮箱 Email",
            "website": "网址 Web",
            "address": "地址 Address",
            "postal": "邮编",
        }
    return {"mobile": "Mobile", "phone": "Office", "fax": "Fax", "email": "Email", "website": "Web", "address": "Address", "postal": "Postal"}


def identity_lines(item: dict) -> list[tuple[str, int, bool]]:
    lines = []
    if item["name"]:
        lines.append((item["name"], 54 if item["layout"] != "竖版" else 48, True))
    if item["jobTitle"]:
        lines.append((item["jobTitle"], 28, False))
    department = item["department"]
    if item["department_en"]:
        department = f"{item['department']} / {item['department_en']}" if item["department"] else item["department_en"]
    if department:
        lines.append((department, 24, False))
    company = item["company"]
    if item["company_en"]:
        company = f"{item['company']}  {item['company_en']}" if item["company"] else item["company_en"]
    if company:
        lines.append((company, 32, True))
    return lines


def contact_lines(item: dict) -> list[str]:
    lab = labels_for(item["lang"])
    rows = []
    for key in ("mobile", "phone", "fax", "website", "address", "postal_code"):
        value = item[key]
        if not value:
            continue
        label_key = {"postal_code": "postal"}.get(key, key)
        rows.append(f"{lab[label_key]}  {value}")
    if item["email"]:
        for email in [part.strip() for part in item["email"].split(";") if part.strip()]:
            rows.append(f"{lab['email']}  {email}")
    return rows


def draw_logo(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str, accent: tuple[int, int, int]):
    draw.ellipse(box, fill=accent)
    face = font(22 if len(text) <= 4 else 16, bold=True)
    bbox = draw.textbbox((0, 0), text, font=face)
    tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
    x = box[0] + (box[2] - box[0] - tw) / 2
    y = box[1] + (box[3] - box[1] - th) / 2 - 2
    draw.text((x, y), text, font=face, fill=(255, 255, 255))


def render_flat(item: dict, scale: float = 1.0) -> Image.Image:
    vertical = item["layout"] == "竖版"
    width, height = (760, 1280) if vertical else (1280, 860)
    theme = THEMES[item["theme"]]
    image = Image.new("RGB", (width, height), theme["bg"])
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 18, height), fill=theme["accent"])
    draw.rectangle((0, 0, width - 1, height - 1), outline=(180, 176, 168), width=2)

    logo_pad = 210 if item["logo_text"] else 0
    if item["logo_text"]:
        draw_logo(draw, (width - 180, 36, width - 48, 168), item["logo_text"], theme["accent"])

    def paint(lines, x, max_width, fill, start_y, gap_extra=8):
        y = start_y
        for text, size, bold in lines:
            face = font(max(16, int(size * scale)), bold=bold)
            for piece in wrap_text(draw, text, face, max_width):
                if y > height - 40:
                    raise RuntimeError(f"{item['card_id']} overflow on {piece}")
                draw.text((x, y), piece, font=face, fill=fill)
                y += int(size * scale) + gap_extra
            y += 4
        return y

    if item["layout"] == "特殊分栏":
        draw.line((width // 2, 36, width // 2, height - 36), fill=(210, 206, 198), width=2)
        y = paint(identity_lines(item), 48, width // 2 - 80, theme["ink"], 48)
        body = font(max(18, int(26 * scale)))
        y2 = 56
        for row in contact_lines(item):
            for piece in wrap_text(draw, row, body, width // 2 - 90):
                draw.text((width // 2 + 36, y2), piece, font=body, fill=theme["muted"])
                y2 += int(34 * scale)
        if y2 > height - 30:
            raise RuntimeError(f"{item['card_id']} column overflow")
        return image

    y = paint(identity_lines(item), 56, width - 90 - logo_pad, theme["ink"], 46, gap_extra=6)
    y += 18
    body = font(max(18, int(26 * scale)))
    max_width = width - 110
    for row in contact_lines(item):
        for piece in wrap_text(draw, row, body, max_width):
            if y > height - 36:
                raise RuntimeError(f"{item['card_id']} overflow on {piece}")
            draw.text((56, y), piece, font=body, fill=theme["muted"])
            y += int(36 * scale)
    return image


def compose_scene(card_image: Image.Image, condition: str) -> Image.Image:
    if condition in {"正面清晰", "小字号"}:
        return card_image

    canvas = Image.new("RGB", (1800, 1250), (214, 206, 194))
    draw = ImageDraw.Draw(canvas)
    angle = 0
    if condition == "倾斜":
        canvas.paste((215, 196, 163), (0, 0, 1800, 1250))
        angle = 13
    elif condition == "深色背景":
        canvas.paste((32, 36, 43), (0, 0, 1800, 1250))
        angle = -7
    elif condition == "浅色背景":
        canvas.paste((243, 241, 236), (0, 0, 1800, 1250))
        angle = 4
    elif condition == "带阴影":
        canvas.paste((236, 231, 223), (0, 0, 1800, 1250))
        angle = 3
    elif condition == "边缘遮挡":
        canvas.paste((198, 184, 154), (0, 0, 1800, 1250))
        angle = -3
    elif condition == "复杂桌面背景":
        canvas.paste((186, 154, 112), (0, 0, 1800, 1250))
        draw.ellipse((70, 80, 280, 290), fill=(92, 63, 46))
        draw.rounded_rectangle((1380, 70, 1720, 360), radius=28, fill=(48, 84, 92))
        draw.rectangle((40, 980, 520, 1180), fill=(230, 222, 206))
        angle = 8
    elif condition == "复杂背景":
        canvas.paste((232, 226, 214), (0, 0, 1800, 1250))
        for i, color in enumerate(((196, 92, 74), (46, 96, 128), (232, 186, 92), (64, 112, 86))):
            draw.rectangle((40 + i * 430, 40, 360 + i * 430, 220), fill=color)
        draw.ellipse((80, 900, 340, 1160), fill=(90, 90, 96))
        draw.rounded_rectangle((1450, 860, 1740, 1160), radius=20, fill=(250, 248, 242))
        angle = -6
    elif condition == "光线不足且文字较小":
        canvas.paste((92, 88, 82), (0, 0, 1800, 1250))
        angle = -2
    elif condition == "模糊低对比":
        canvas.paste((176, 176, 172), (0, 0, 1800, 1250))
        angle = 2
    elif condition == "倾斜加阴影加小字":
        canvas.paste((168, 150, 124), (0, 0, 1800, 1250))
        angle = 11
    elif condition == "复杂Logo":
        return card_image
    elif condition == "复杂桌面背景" and False:
        pass

    rotated = card_image.rotate(angle, expand=True, resample=Image.Resampling.BICUBIC, fillcolor=(0, 0, 0, 0))
    # rotate of RGB has no alpha; use a mask by comparing to a throwaway fill via RGBA
    rgba = card_image.convert("RGBA")
    rotated = rgba.rotate(angle, expand=True, resample=Image.Resampling.BICUBIC)
    x = (canvas.width - rotated.width) // 2
    y = (canvas.height - rotated.height) // 2
    shadow = Image.new("RGBA", rotated.size, (0, 0, 0, 0))
    shadow_draw = ImageDraw.Draw(shadow)
    shadow_draw.bitmap((18, 22), rotated.split()[-1], fill=(0, 0, 0, 90))
    canvas.paste(shadow, (x, y), shadow)
    canvas.paste(rotated, (x, y), rotated)

    if condition == "边缘遮挡":
        cover = ImageDraw.Draw(canvas)
        cover.polygon([(1320, 860), (1760, 980), (1760, 1220), (1180, 1220)], fill=(126, 92, 58))

    if condition == "光线不足且文字较小":
        canvas = ImageEnhance.Brightness(canvas).enhance(0.62)
        canvas = ImageEnhance.Contrast(canvas).enhance(0.8)
    elif condition == "模糊低对比":
        canvas = canvas.filter(ImageFilter.GaussianBlur(radius=1.15))
        canvas = ImageEnhance.Contrast(canvas).enhance(0.72)
    elif condition == "倾斜加阴影加小字":
        canvas = ImageEnhance.Brightness(canvas).enhance(0.78)
        canvas = ImageEnhance.Contrast(canvas).enhance(0.85)
    return canvas.convert("RGB")


def validate(cards: list[dict]) -> None:
    if len(cards) != 40:
        raise SystemExit(f"expected 40 cards, got {len(cards)}")
    ids = [c["card_id"] for c in cards]
    if ids != [f"C{i:03d}" for i in range(1, 41)]:
        raise SystemExit("card ids are not C001–C040 in order")
    mobiles = [c["mobile"] for c in cards if c["mobile"]]
    phones = [c["phone"] for c in cards if c["phone"]]
    if len(mobiles) != len(set(mobiles)):
        raise SystemExit("duplicate mobile")
    if len(phones) != len(set(phones)):
        raise SystemExit("duplicate office phone")
    for c in cards:
        for email in [p.strip() for p in c["email"].split(";") if p.strip()]:
            if not (email.endswith("@example.com") or email.endswith(".example")):
                raise SystemExit(f"non-reserved email on {c['card_id']}: {email}")
        for key in ("mobile", "phone", "fax"):
            if c[key] and not (c[key].startswith("+1-555-") or c[key].startswith("+86 ")):
                raise SystemExit(f"unexpected phone on {c['card_id']}: {c[key]}")
    by_id = {c["card_id"]: c for c in cards}
    if by_id["C007"]["website"] or by_id["C025"]["website"] or by_id["C034"]["website"]:
        raise SystemExit("designated missing-website cards have a website")
    if not (by_id["C025"]["mobile"] and by_id["C025"]["phone"] and by_id["C025"]["fax"]):
        raise SystemExit("C025 must have three phone roles")
    if by_id["C030"]["name"] or by_id["C032"]["company"] or by_id["C034"]["company"]:
        raise SystemExit("missing-field cards are not actually missing")
    if "Apple" not in by_id["C034"]["address"]:
        raise SystemExit("C034 address lost Apple Street")


def write_xlsx(cards: list[dict]) -> None:
    wb = Workbook()
    intro = wb.active
    intro.title = "使用说明"
    intro["A1"] = "测试名片 Ground Truth"
    intro["A1"].font = Font(name="微软雅黑", bold=True, size=16, color="1F4E79")
    lines = [
        "本表是学生自制的虚构测试集，共 40 张（作业建议 30～50 张）。",
        "作业原文只写了“建议教师提供统一测试名片”，这一份作业文件里没有附带官方图片或官方答案。",
        "如果教师之后发放固定测试集，用官方数据替换本表，不要用这些虚构名片覆盖官方集。",
        "所有姓名、公司、电话、邮箱和地址都是虚构的。邮箱只用 example.com 和 .example。电话只用 +1-555-01xx 或 +86 的占位号。",
        "空白单元格的意思是：这张名片上没有该字段。识别结果必须为空，不允许补全。",
        "email 列如果有两个地址，用英文分号加空格分隔，例如 a@example.com; b@example.com。两个都要保留。",
        "company 只记录一个机构。名片上同时印了中文名和英文名时，英文名是同一机构，不是第二家公司，说明写在 notes。",
        "image_file 是相对于项目文件夹的路径。图片里只印名片上真实存在的文字，没有把 card_id 印进画面。",
        "拍摄条件已经做进图片：倾斜、深色或浅色背景、阴影、边缘遮挡、小字、模糊、复杂桌面。这是模拟图，不是从参考 App 截来的。",
        "Logo 上的字母不是姓名或公司，除非 notes 明确说它就是公司名。",
    ]
    for index, line in enumerate(lines, start=3):
        intro.cell(index, 1, line).font = Font(name="微软雅黑", size=12)
        intro.cell(index, 1).alignment = Alignment(wrap_text=True, vertical="center")
        intro.row_dimensions[index].height = 22
    intro.column_dimensions["A"].width = 120

    headers = [
        "card_id", "image_file", "category", "layout", "capture_condition",
        "name", "company", "department", "jobTitle", "mobile", "phone", "fax",
        "email", "website", "address", "postal_code", "notes",
    ]
    ws = wb.create_sheet("测试名片", 0)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
    thin = Border(
        left=Side(style="thin", color="D9E2EF"),
        right=Side(style="thin", color="D9E2EF"),
        top=Side(style="thin", color="D9E2EF"),
        bottom=Side(style="thin", color="D9E2EF"),
    )
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(1, col, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    wraps = {"notes", "address", "email", "company"}
    for row_index, item in enumerate(cards, start=2):
        values = {
            "card_id": item["card_id"],
            "image_file": f"test_card_images/{item['card_id']}.png",
            "category": item["category"],
            "layout": item["layout"],
            "capture_condition": item["capture_condition"],
            "name": item["name"],
            "company": item["company"],
            "department": item["department"],
            "jobTitle": item["jobTitle"],
            "mobile": item["mobile"],
            "phone": item["phone"],
            "fax": item["fax"],
            "email": item["email"],
            "website": item["website"],
            "address": item["address"],
            "postal_code": item["postal_code"],
            "notes": item["notes"],
        }
        for col, header in enumerate(headers, start=1):
            value = values[header] or None
            cell = ws.cell(row_index, col, value)
            cell.font = Font(name="微软雅黑", size=10)
            cell.alignment = Alignment(wrap_text=header in wraps, vertical="center")
            cell.border = thin
            if item["website"] == "" and header == "website":
                cell.fill = PatternFill("solid", fgColor="FFF4E5")
        ws.row_dimensions[row_index].height = 36
    ws.auto_filter.ref = f"A1:Q{len(cards) + 1}"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(cards) + 1}"
    widths = {
        "A": 12, "B": 28, "C": 18, "D": 14, "E": 22, "F": 16, "G": 22, "H": 16,
        "I": 22, "J": 20, "K": 20, "L": 20, "M": 46, "N": 32, "O": 28, "P": 14, "Q": 72,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.row_dimensions[1].height = 24

    link = wb.create_sheet("用例对照")
    link["A1"] = "这些对照是本虚构测试集的用法，不是参考 App 的观察结果。"
    link["A1"].font = Font(name="微软雅黑", bold=True, size=12)
    pairs = [
        ("T06 边界检测", "C002、C009、C017，以及 C006 的边缘遮挡"),
        ("T07 OCR 原文", "先用 C001，再扩展到 C008、C015"),
        ("T08 字段分类", "C001 英文齐全；C008 中文齐全；C015 中英混排"),
        ("T09 不得虚构网站", "C007 是主样例。C010、C011、C015、C016、C019、C021、C024、C025、C028、C029、C030、C031、C032、C033、C034、C037 也没有网站"),
        ("T10 多电话", "C025 是 Mobile / Office / Fax 主样例；C026、C014 是中文三号码"),
        ("多邮箱", "C027、C028"),
        ("字段缺失", "C029 无职位无网站；C030 无姓名；C031 无职位；C032 只有姓名和手机"),
        ("字段歧义", "C015 中英职称；C022 姓与公司同词；C033 Jordan Lee；C034 Apple Street；C040 柏年 / 柏年实验室"),
        ("竖版或特殊布局", "C021、C022、C023、C024"),
        ("较差画质", "C012 小字；C035 光线不足；C036 模糊；C037 倾斜加阴影加小字"),
        ("复杂 Logo 或背景", "C014、C038、C039、C040"),
    ]
    link["A3"] = "用途"
    link["B3"] = "名片"
    for col in (1, 2):
        link.cell(3, col).font = Font(name="微软雅黑", bold=True, color="FFFFFF")
        link.cell(3, col).fill = header_fill
    for offset, (use, cards_text) in enumerate(pairs, start=4):
        link.cell(offset, 1, use).font = Font(name="微软雅黑", size=11)
        link.cell(offset, 2, cards_text).font = Font(name="微软雅黑", size=11)
        link.cell(offset, 2).alignment = Alignment(wrap_text=True)
        link.row_dimensions[offset].height = 32
    link.column_dimensions["A"].width = 28
    link.column_dimensions["B"].width = 110

    wb.save(XLSX_PATH)


def main() -> None:
    validate(CARDS)
    os.makedirs(IMAGE_DIR, exist_ok=True)
    for item in CARDS:
        scale = 0.78 if item["capture_condition"] in {"小字号", "光线不足且文字较小", "倾斜加阴影加小字"} else 1.0
        flat = render_flat(item, scale=scale)
        final = compose_scene(flat, item["capture_condition"])
        path = os.path.join(IMAGE_DIR, f"{item['card_id']}.png")
        final.save(path, "PNG")
        print(f"{item['card_id']} {final.size[0]}x{final.size[1]}")
    write_xlsx(CARDS)
    print(f"wrote {XLSX_PATH}")


if __name__ == "__main__":
    main()
